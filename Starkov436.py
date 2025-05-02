import openpyxl
from openpyxl.styles.fills import PatternFill

work_book_observed = openpyxl.load_workbook(filename="Observed.xlsx")                               
sheet_obs = work_book_observed.active                                                       
    
work_book_expected = openpyxl.load_workbook(filename="Expected.xlsx")                                
sheet_exp = work_book_expected.active                                                       

work_book_result = openpyxl.load_workbook(filename="Observed-Expected.xlsx")                         
sheet_res = work_book_result.active      

def compare_cells ():                                      
    for i in range(1, sheet_obs.max_row + 1):                                                   
        for k in range(1, sheet_obs.max_column + 1):                                            
            if k == 1 or i == 1 or k == sheet_obs.max_column or i == sheet_obs.max_row:         
                sheet_res.cell(i, k).value = sheet_obs.cell(i, k).value                         
            else:                                                                               
                val = sheet_obs.cell(i, k).value - sheet_exp.cell(i, k).value                   
                if abs(val) > 10:                                                               
                    fill_color = PatternFill(fgColor = 'e5eff7', fill_type='solid')             
                    sheet_res.cell(i, k).fill = fill_color                                      
                sheet_res.cell(i, k).value = val 

compare_cells()            

work_book_result.save('Observed-Expected.xlsx')


